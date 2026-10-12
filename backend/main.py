import asyncio
import base64
import hashlib
import hmac
import io
import json
import logging
import os
import re
import secrets
import sqlite3
import threading
import time
import uuid
from contextlib import contextmanager
from typing import Any, Dict, List, Optional
from urllib.parse import quote

import httpx
import fitz
import pdfplumber
import uvicorn
try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
    from psycopg2.pool import ThreadedConnectionPool
except ImportError:  # pragma: no cover - SQLite fallback still works without Postgres dependency.
    psycopg2 = None
    RealDictCursor = None
    ThreadedConnectionPool = None
from docx import Document
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, validator
from pypdf import PdfReader

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))
load_dotenv(".env")
LOGGER = logging.getLogger(__name__)

DEFAULT_API_BASE_URL = "https://api.siliconflow.cn/v1"
DEFAULT_MODEL = "deepseek-ai/DeepSeek-V4-Flash"
DEFAULT_OCR_MODEL = "Qwen/Qwen3-VL-8B-Instruct"
DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
DB_PATH = os.path.join(DATA_DIR, "eviden.db")
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if not DATABASE_URL and os.getenv("SUPABASE_DB_PASSWORD"):
    db_password = quote(os.getenv("SUPABASE_DB_PASSWORD", ""), safe="")
    db_host = os.getenv("SUPABASE_DB_HOST", "db.fjuiloufsrvewjlvijys.supabase.co")
    db_user = quote(os.getenv("SUPABASE_DB_USER", "postgres"), safe="")
    DATABASE_URL = f"postgresql://{db_user}:{db_password}@{db_host}:5432/postgres"
TOKEN_TTL_SECONDS = 60 * 60 * 24 * 30
_DB_POOL = None
_ANALYSIS_JOBS: Dict[str, Dict[str, Any]] = {}
_ANALYSIS_JOBS_LOCK = threading.Lock()
DIMENSION_WEIGHTS = {
    "核心能力匹配": 30,
    "经历证据强度": 25,
    "业务场景匹配": 15,
    "结果量化程度": 10,
    "角色责任匹配": 10,
    "关键词与表达": 10,
}


class AnalyzeRequest(BaseModel):
    resume: str = Field(..., min_length=80, max_length=50000)
    jd: str = Field(..., min_length=50, max_length=30000)
    target: str = Field(..., min_length=2, max_length=120)
    company: Optional[str] = Field("", max_length=120)
    business_unit: Optional[str] = Field("", max_length=120)

    @validator("resume", "jd", "target", "company", "business_unit", pre=True, always=True)
    def strip_text(cls, value: Optional[str]) -> str:
        return str(value or "").strip()


class ResumeRewriteRequest(BaseModel):
    resume: str = Field(..., min_length=80, max_length=50000)
    target: str = Field(..., min_length=2, max_length=120)
    jd: Optional[str] = Field("", max_length=30000)
    company: Optional[str] = Field("", max_length=120)
    business_unit: Optional[str] = Field("", max_length=120)

    @validator("resume", "target", "jd", "company", "business_unit", pre=True, always=True)
    def strip_text(cls, value: Optional[str]) -> str:
        return str(value or "").strip()


class RewriteBullet(BaseModel):
    section: str
    before: str
    after: str
    reason: str
    evidence_boundary: str


class ResumeRewriteResult(BaseModel):
    target: str
    positioning: str
    strategy: str
    bullets: List[RewriteBullet]
    gaps: List[str]
    hr_intro: str
    cautions: List[str]
    final_resume: str
    model: str


class ResumeRewriteJobStatus(BaseModel):
    job_id: str
    status: str
    result: Optional[ResumeRewriteResult] = None
    error: str = ""
    created_at: int
    updated_at: int


class AuthRequest(BaseModel):
    phone: str = Field(..., min_length=6, max_length=20)
    password: str = Field(..., min_length=6, max_length=72)

    @validator("phone")
    def validate_phone(cls, value: str) -> str:
        cleaned = re.sub(r"\D", "", value)
        if len(cleaned) < 6 or len(cleaned) > 20:
            raise ValueError("请输入有效手机号")
        return cleaned


class UserInfo(BaseModel):
    id: int
    phone: str
    created_at: int


class AuthResponse(BaseModel):
    token: str
    user: UserInfo


class CloudReport(BaseModel):
    id: str
    record: Dict[str, Any]
    updated_at: int


class CloudReportRequest(BaseModel):
    id: str = Field(..., min_length=3, max_length=120)
    record: Dict[str, Any]


class CloudReportsResponse(BaseModel):
    reports: List[CloudReport]


class DefaultResumeRequest(BaseModel):
    filename: str = Field(..., min_length=1, max_length=255)
    resume_text: str = Field(..., min_length=80, max_length=50000)

    @validator("filename", pre=True)
    def strip_filename(cls, value: str) -> str:
        return str(value or "").strip()

    @validator("resume_text", pre=True)
    def normalize_resume_text(cls, value: str) -> str:
        normalized = str(value or "").strip()
        normalized = re.sub(r"\r\n?", "\n", normalized)
        normalized = re.sub(r"[ \t]+", " ", normalized)
        return re.sub(r"\n{3,}", "\n\n", normalized)


class DefaultResume(BaseModel):
    filename: str
    resume_text: str
    char_count: int
    preview: str
    updated_at: int


class DefaultResumeResponse(BaseModel):
    resume: Optional[DefaultResume] = None


class DeleteResponse(BaseModel):
    ok: bool


class Dimension(BaseModel):
    name: str
    weight: int = 0
    value: int = Field(..., ge=0, le=100)
    note: str
    points: float = 0
    bonus_evidence: List[str] = Field(default_factory=list)
    deduction_reason: str = ""
    remedy: str = ""


class Evidence(BaseModel):
    req: str
    priority: str
    jdQuote: str
    proof: str
    resumeQuote: str
    level: str
    confidence: int = Field(..., ge=0, le=100)
    diagnosis: str
    suggestion: str = ""

    @validator("level")
    def validate_level(cls, value: str) -> str:
        mapping = {
            "强": "直接证据",
            "中": "可迁移证据",
            "弱": "可迁移证据",
            "直接证据": "直接证据",
            "可迁移证据": "可迁移证据",
            "无相关证据": "无相关证据",
        }
        return mapping.get(value, "可迁移证据")


class Gap(BaseModel):
    level: str
    title: str
    desc: str


class Action(BaseModel):
    title: str
    desc: str
    effort: str


class Rewrite(BaseModel):
    before: str
    after: str


class AnalysisPayload(BaseModel):
    dimensions: List[Dimension]
    summary: str
    evidence: List[Evidence]
    gaps: List[Gap]
    actions: List[Action]
    rewrites: List[Rewrite]
    questions: List[str]
    score_model: str = "通用证据链评分模型 v1"
    position_type: str = "通用岗位"
    confidence_label: str = "中"
    confidence_reason: str = "基于简历与 JD 的文本完整度、证据数量和引用可核验程度判断"


class AnalysisResult(AnalysisPayload):
    target: str
    resumeName: str
    jdName: str
    score: int
    priority: str
    action: str
    model: str


class AnalysisJobCreated(BaseModel):
    job_id: str
    status: str


class AnalysisJobStatus(BaseModel):
    job_id: str
    status: str
    result: Optional[AnalysisResult] = None
    error: str = ""
    created_at: int
    updated_at: int


class ScreeningJD(BaseModel):
    title: str = Field(..., min_length=2, max_length=120)
    jd: str = Field(..., min_length=50, max_length=30000)
    company: str = Field("", max_length=120)
    business_unit: str = Field("", max_length=120)


class ScreeningRequest(BaseModel):
    resume: str = Field(..., min_length=80, max_length=50000)
    jds: List[ScreeningJD]

    @validator("jds")
    def validate_jd_count(cls, value: List[ScreeningJD]) -> List[ScreeningJD]:
        if len(value) < 2 or len(value) > 9:
            raise ValueError("请添加 2–9 个岗位")
        return value


class ScreeningMatch(BaseModel):
    jd_index: int
    title: str
    company: str = ""
    score: int = Field(..., ge=0, le=100)
    priority: str
    verdict: str
    match_points: List[str]
    gaps: List[str]


class ScreeningResponse(BaseModel):
    results: List[ScreeningMatch]
    model: str


class HealthResponse(BaseModel):
    ok: bool
    model: str
    configured: bool
    db_backend: str = "sqlite"


class ResumeParseResponse(BaseModel):
    filename: str
    text: str
    char_count: int
    preview: str


class ResumeParseJsonRequest(BaseModel):
    filename: str = Field(..., min_length=1, max_length=255)
    content_base64: str = Field(..., min_length=1)


def database_url() -> str:
    explicit_url = DATABASE_URL or read_local_secret("DATABASE_URL") or read_config_module("DATABASE_URL")
    if explicit_url:
        return explicit_url.strip()
    db_password = get_secret("SUPABASE_DB_PASSWORD")
    if not db_password:
        return ""
    db_host = get_secret("SUPABASE_DB_HOST") or "db.fjuiloufsrvewjlvijys.supabase.co"
    db_user = get_secret("SUPABASE_DB_USER") or "postgres"
    return "postgresql://{0}:{1}@{2}:5432/postgres".format(quote(db_user, safe=""), quote(db_password, safe=""), db_host)


def use_postgres() -> bool:
    return bool(database_url())


@contextmanager
def db_connection():  # type: ignore[no-untyped-def]
    global _DB_POOL
    if use_postgres():
        if psycopg2 is None or ThreadedConnectionPool is None:
            raise RuntimeError("Postgres dependency is not installed")
        if _DB_POOL is None:
            _DB_POOL = ThreadedConnectionPool(
                1,
                6,
                database_url(),
                cursor_factory=RealDictCursor,
                connect_timeout=8,
                keepalives=1,
                keepalives_idle=10,
                keepalives_interval=5,
                keepalives_count=3,
                sslmode="require",
                options="-c statement_timeout=15000",
            )
        connection = _DB_POOL.getconn()
        try:
            try:
                with connection.cursor() as health_cursor:
                    health_cursor.execute("SELECT 1")
                    health_cursor.fetchone()
            except Exception:
                _DB_POOL.putconn(connection, close=True)
                connection = _DB_POOL.getconn()
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            _DB_POOL.putconn(connection)
        return
    os.makedirs(DATA_DIR, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def sql_params(sql: str) -> str:
    return sql.replace("?", "%s") if use_postgres() else sql


def fetch_one(sql: str, params: tuple = ()):  # type: ignore[no-untyped-def]
    with db_connection() as connection:
        if use_postgres():
            with connection.cursor() as cursor:
                cursor.execute(sql_params(sql), params)
                return cursor.fetchone()
        return connection.execute(sql, params).fetchone()


def fetch_all(sql: str, params: tuple = ()):  # type: ignore[no-untyped-def]
    with db_connection() as connection:
        if use_postgres():
            with connection.cursor() as cursor:
                cursor.execute(sql_params(sql), params)
                return cursor.fetchall()
        return connection.execute(sql, params).fetchall()


def execute_sql(sql: str, params: tuple = ()) -> None:
    with db_connection() as connection:
        if use_postgres():
            with connection.cursor() as cursor:
                cursor.execute(sql_params(sql), params)
            return
        connection.execute(sql, params)


def init_db() -> None:
    if use_postgres():
        statements = [
            """
            CREATE TABLE IF NOT EXISTS users (
                id SERIAL PRIMARY KEY,
                phone TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at INTEGER NOT NULL
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                expires_at INTEGER NOT NULL,
                created_at INTEGER NOT NULL
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS reports (
                id TEXT NOT NULL,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                record_json TEXT NOT NULL,
                updated_at INTEGER NOT NULL,
                PRIMARY KEY(id, user_id)
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS analysis_jobs (
                id TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                status TEXT NOT NULL,
                request_json TEXT NOT NULL,
                result_json TEXT,
                error TEXT NOT NULL DEFAULT '',
                created_at INTEGER NOT NULL,
                updated_at INTEGER NOT NULL
            )
            """,
            """
            CREATE TABLE IF NOT EXISTS user_resumes (
                user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
                filename TEXT NOT NULL,
                resume_text TEXT NOT NULL,
                char_count INTEGER NOT NULL,
                preview TEXT NOT NULL,
                updated_at INTEGER NOT NULL
            )
            """,
        ]
        with db_connection() as connection:
            with connection.cursor() as cursor:
                for statement in statements:
                    cursor.execute(statement)
        return

    with db_connection() as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                phone TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                created_at INTEGER NOT NULL
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL,
                expires_at INTEGER NOT NULL,
                created_at INTEGER NOT NULL,
                FOREIGN KEY(user_id) REFERENCES users(id)
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS reports (
                id TEXT NOT NULL,
                user_id INTEGER NOT NULL,
                record_json TEXT NOT NULL,
                updated_at INTEGER NOT NULL,
                PRIMARY KEY(id, user_id),
                FOREIGN KEY(user_id) REFERENCES users(id)
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS analysis_jobs (
                id TEXT PRIMARY KEY,
                user_id INTEGER NOT NULL,
                status TEXT NOT NULL,
                request_json TEXT NOT NULL,
                result_json TEXT,
                error TEXT NOT NULL DEFAULT '',
                created_at INTEGER NOT NULL,
                updated_at INTEGER NOT NULL,
                FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS user_resumes (
                user_id INTEGER PRIMARY KEY,
                filename TEXT NOT NULL,
                resume_text TEXT NOT NULL,
                char_count INTEGER NOT NULL,
                preview TEXT NOT NULL,
                updated_at INTEGER NOT NULL,
                FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
            )
            """
        )


def is_integrity_error(error: Exception) -> bool:
    if isinstance(error, sqlite3.IntegrityError):
        return True
    return bool(psycopg2 is not None and isinstance(error, psycopg2.IntegrityError))


def create_user(phone: str, password_hash: str, created_at: int) -> int:
    with db_connection() as connection:
        if use_postgres():
            with connection.cursor() as cursor:
                cursor.execute(
                    "INSERT INTO users(phone, password_hash, created_at) VALUES (%s, %s, %s) RETURNING id",
                    (phone, password_hash, created_at),
                )
                row = cursor.fetchone()
                return int(row["id"])
        cursor = connection.execute(
            "INSERT INTO users(phone, password_hash, created_at) VALUES (?, ?, ?)",
            (phone, password_hash, created_at),
        )
        return int(cursor.lastrowid)


def hash_password(password: str, salt: Optional[str] = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 180000)
    return "{0}${1}".format(salt, digest.hex())


def verify_password(password: str, stored_hash: str) -> bool:
    try:
        salt, _ = stored_hash.split("$", 1)
    except ValueError:
        return False
    return hmac.compare_digest(hash_password(password, salt), stored_hash)


def create_session(user_id: int) -> str:
    token = secrets.token_urlsafe(32)
    now = int(time.time())
    execute_sql(
        "INSERT INTO sessions(token, user_id, expires_at, created_at) VALUES (?, ?, ?, ?)",
        (token, user_id, now + TOKEN_TTL_SECONDS, now),
    )
    return token


def authenticate_and_create_session(phone: str, password: str):  # type: ignore[no-untyped-def]
    token = secrets.token_urlsafe(32)
    now = int(time.time())
    with db_connection() as connection:
        if use_postgres():
            with connection.cursor() as cursor:
                cursor.execute("SELECT * FROM users WHERE phone = %s", (phone,))
                row = cursor.fetchone()
                if row is None or not verify_password(password, str(row["password_hash"])):
                    return None, ""
                cursor.execute(
                    "INSERT INTO sessions(token, user_id, expires_at, created_at) VALUES (%s, %s, %s, %s)",
                    (token, int(row["id"]), now + TOKEN_TTL_SECONDS, now),
                )
                return row, token
        row = connection.execute("SELECT * FROM users WHERE phone = ?", (phone,)).fetchone()
        if row is None or not verify_password(password, str(row["password_hash"])):
            return None, ""
        connection.execute(
            "INSERT INTO sessions(token, user_id, expires_at, created_at) VALUES (?, ?, ?, ?)",
            (token, int(row["id"]), now + TOKEN_TTL_SECONDS, now),
        )
        return row, token


def user_from_row(row) -> UserInfo:  # type: ignore[no-untyped-def]
    return UserInfo(id=int(row["id"]), phone=str(row["phone"]), created_at=int(row["created_at"]))


def default_resume_from_row(row) -> DefaultResume:  # type: ignore[no-untyped-def]
    return DefaultResume(
        filename=str(row["filename"]),
        resume_text=str(row["resume_text"]),
        char_count=int(row["char_count"]),
        preview=str(row["preview"]),
        updated_at=int(row["updated_at"]),
    )


def current_user(authorization: Optional[str] = Header(default=None)) -> UserInfo:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="请先登录")
    token = authorization.split(" ", 1)[1].strip()
    now = int(time.time())
    row = fetch_one(
        """
        SELECT users.id, users.phone, users.created_at
        FROM sessions
        JOIN users ON users.id = sessions.user_id
        WHERE sessions.token = ? AND sessions.expires_at > ?
        """,
        (token, now),
    )
    if row is None:
        raise HTTPException(status_code=401, detail="登录已过期，请重新登录")
    return user_from_row(row)


def clean_resume_text(text: str) -> str:
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf_resume_with_pypdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    chunks = []
    for page in reader.pages:
        chunks.append(page.extract_text() or "")
    return clean_resume_text("\n".join(chunks))


def parse_pdf_resume_with_fitz(content: bytes) -> str:
    document = fitz.open(stream=content, filetype="pdf")
    chunks = []
    for page in document:
        blocks = page.get_text("blocks")
        blocks = sorted(blocks, key=lambda block: (round(block[1] / 10), block[0]))
        page_text = "\n".join(block[4].strip() for block in blocks if block[4].strip())
        chunks.append(page_text)
    document.close()
    return clean_resume_text("\n\n".join(chunks))


def parse_pdf_resume_with_pdfplumber(content: bytes) -> str:
    chunks = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=1.5, y_tolerance=3) or ""
            tables = []
            for table in page.extract_tables() or []:
                rows = []
                for row in table:
                    cells = [str(cell or "").strip() for cell in row if str(cell or "").strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    tables.append("\n".join(rows))
            chunks.append("\n".join(part for part in [text, "\n".join(tables)] if part.strip()))
    return clean_resume_text("\n\n".join(chunks))


def score_extracted_text(text: str) -> float:
    if not text:
        return 0
    replacement_penalty = text.count("�") + text.count("·") * 0.4
    useful_chars = sum(1 for char in text if char.isalnum() or "\u4e00" <= char <= "\u9fff")
    line_bonus = min(len([line for line in text.splitlines() if line.strip()]), 80) * 3
    contact_bonus = 20 if re.search(r"@|\b1[3-9]\d{9}\b|linkedin|github", text, re.I) else 0
    section_bonus = 15 * len(re.findall(r"教育|经历|项目|技能|工作|实习|experience|education|project|skills", text, re.I))
    return useful_chars + line_bonus + contact_bonus + section_bonus - replacement_penalty * 8


def best_pdf_text(content: bytes) -> str:
    candidates = []
    try:
        fast_text = parse_pdf_resume_with_fitz(content)
        if fast_text:
            candidates.append(fast_text)
            if len(fast_text) >= 160 and score_extracted_text(fast_text) >= 110:
                return fast_text
    except Exception as error:
        LOGGER.info("PDF parser %s failed: %s", parse_pdf_resume_with_fitz.__name__, error)
    for parser in (parse_pdf_resume_with_pdfplumber, parse_pdf_resume_with_pypdf):
        try:
            parsed = parser(content)
            if parsed:
                candidates.append(parsed)
        except Exception as error:
            LOGGER.info("PDF parser %s failed: %s", parser.__name__, error)
    return max(candidates, key=score_extracted_text) if candidates else ""


def parse_pdf_resume(content: bytes) -> str:
    text = best_pdf_text(content)
    if len(text) < 80:
        raise HTTPException(status_code=422, detail="PDF 可提取文本过少，可能是扫描版或图片版简历；请换 DOCX 或可复制文本的 PDF。")
    return text


def render_pdf_pages_for_ocr(content: bytes, max_pages: int = 4) -> List[bytes]:
    document = fitz.open(stream=content, filetype="pdf")
    images = []
    try:
        for page_index in range(min(len(document), max_pages)):
            page = document[page_index]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            images.append(pixmap.tobytes("png"))
    finally:
        document.close()
    return images


async def request_ocr_for_image(image: bytes, page_number: int) -> str:
    api_key = get_secret("SILICONFLOW_API_KEY")
    if not api_key:
        raise HTTPException(status_code=500, detail="OCR 服务未配置，请稍后再试。")
    image_base64 = base64.b64encode(image).decode("utf-8")
    payload = {
        "model": ocr_model_name(),
        "messages": [
            {
                "role": "user",
                "content": [
                    {
                        "type": "image_url",
                        "image_url": {"url": "data:image/png;base64,{0}".format(image_base64)},
                    },
                    {
                        "type": "text",
                        "text": (
                            "请从这张简历 PDF 第 {0} 页图片中提取全部可读文字。"
                            "按从上到下、从左到右的阅读顺序输出，保留姓名、联系方式、教育、经历、项目、技能、时间、公司、职位、项目符号和表格信息。"
                            "不要总结，不要翻译，不要添加解释；如果没有可读文字，只输出空字符串。"
                        ).format(page_number),
                    },
                ],
            }
        ],
        "temperature": 0,
        "max_tokens": 2400,
    }
    headers = {"Authorization": "Bearer {0}".format(api_key), "Content-Type": "application/json"}
    timeout = httpx.Timeout(90.0, connect=15.0)
    try:
        async with httpx.AsyncClient(timeout=timeout, trust_env=False) as client:
            response = await client.post(
                "{0}/chat/completions".format(api_base_url().rstrip("/")),
                headers=headers,
                json=payload,
            )
            if response.status_code >= 400:
                LOGGER.warning("OCR request failed: %s %s", response.status_code, response.text[:300])
                raise HTTPException(status_code=502, detail="扫描版 PDF OCR 失败，请稍后重试。")
            data = response.json()
    except HTTPException:
        raise
    except Exception as error:
        LOGGER.warning("OCR request error: %s", error)
        raise HTTPException(status_code=502, detail="扫描版 PDF OCR 失败，请稍后重试。") from error
    return clean_resume_text(str(data.get("choices", [{}])[0].get("message", {}).get("content", "")))


async def parse_pdf_resume_with_ocr_fallback(content: bytes) -> str:
    text = best_pdf_text(content)
    if len(text) >= 120 and score_extracted_text(text) >= 100:
        return text
    images = render_pdf_pages_for_ocr(content)
    if not images:
        raise HTTPException(status_code=422, detail="PDF 解析失败，请换一个文件重试。")
    page_texts = []
    for index, image in enumerate(images, start=1):
        page_text = await request_ocr_for_image(image, index)
        if page_text:
            page_texts.append(page_text)
    ocr_text = clean_resume_text("\n\n".join(page_texts))
    if len(ocr_text) < 80:
        raise HTTPException(status_code=422, detail="OCR 识别到的简历内容太少，请确认文件清晰或换 DOCX 重试。")
    return ocr_text


def parse_docx_resume(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return clean_resume_text("\n".join(chunks))


def parse_txt_resume(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return clean_resume_text(content.decode(encoding))
        except UnicodeDecodeError:
            continue
    return clean_resume_text(content.decode("utf-8", errors="ignore"))


def parse_resume_file(filename: str, content: bytes) -> str:
    lowered = filename.lower()
    if lowered.endswith(".pdf"):
        return parse_pdf_resume(content)
    if lowered.endswith(".docx"):
        return parse_docx_resume(content)
    if lowered.endswith(".txt"):
        return parse_txt_resume(content)
    raise HTTPException(status_code=400, detail="目前仅支持 PDF、DOCX 或 TXT 简历文件。")


def read_local_secret(name: str) -> str:
    current_dir = os.path.dirname(__file__)
    candidate_paths = [
        os.path.join(current_dir, ".env"),
        os.path.join(os.getcwd(), ".env"),
    ]
    for path in candidate_paths:
        try:
            with open(path, "r") as file_obj:
                for line in file_obj:
                    if line.startswith(name + "="):
                        return line.split("=", 1)[1].strip().strip('"').strip("'")
        except OSError:
            continue
    return ""


def read_config_module(name: str) -> str:
    try:
        import config_secret  # type: ignore
        return str(getattr(config_secret, name, "")).strip()
    except Exception:
        return ""


def get_secret(name: str) -> str:
    return os.environ.get(name, "").strip() or read_local_secret(name) or read_config_module(name)


def api_base_url() -> str:
    return os.environ.get("SILICONFLOW_API_BASE_URL", "").strip() or read_local_secret("SILICONFLOW_API_BASE_URL") or read_config_module("SILICONFLOW_API_BASE_URL") or DEFAULT_API_BASE_URL


def model_name() -> str:
    return (os.environ.get("SILICONFLOW_MODEL", "").strip() or read_local_secret("SILICONFLOW_MODEL") or read_config_module("SILICONFLOW_MODEL") or DEFAULT_MODEL)


def ocr_model_name() -> str:
    return (os.environ.get("SILICONFLOW_OCR_MODEL", "").strip() or read_local_secret("SILICONFLOW_OCR_MODEL") or read_config_module("SILICONFLOW_OCR_MODEL") or DEFAULT_OCR_MODEL)


def build_system_prompt() -> str:
    return """你是 Eviden 的岗位匹配分析引擎。你的任务是比较一份简历与一个目标 JD，生成可核验的求职申请策略。

必须遵守：
1. 只能使用用户提供的简历与 JD，不得编造经历、数据、职责、工具或结果。
2. jdQuote 必须引用 JD 中真实存在的片段；resumeQuote 必须引用简历中真实存在的最短完整片段。语义相关但场景不完全一致的经历应标记为“可迁移证据”，不要因为不是同岗位或措辞不同就写成无证据。
3. 按 /job-match 方法拆解：先识别硬性门槛，再看核心职责、能力要求、加分项和申请约束。公司与业务方向用于理解行业、产品形态、客户类型、商业模式和工作场景，并帮助判断 JD 中同一能力在具体岗位场景下的含义；这些信息可以影响证据、缺口、行动、改写和面试问题的优先级，但不得作为简历证据、独立加分项或单独评分维度，也无需在报告中额外声明。
4. 每项要求的状态必须在诊断中明确区分：已匹配、表达缺口、证据不足、真实缺口、待确认。
5. 区分三类问题：真实能力缺口、已有能力但证据不足、仅表达方式不佳。
6. 评分应保守；关键词相同不等于具备强证据；硬性门槛缺失时不得被其他优势完全抵消。
7. 对完全相同的简历、JD、目标岗位、公司和事业部输入，必须保持一致的维度分、总分、优先级和结论；不要因为表达随机性改变评分。
8. 使用通用证据链评分模型 v1，不按岗位类型改变权重，但可以在 position_type 中识别岗位大类，帮助解释语境。
9. 每个维度必须给出加分证据、扣分原因和补救建议，让用户能判断分数是否合理。
10. 改写建议必须遵循 ASu-skills 中 /great-resume 的原则：只做事实重组、岗位化表达和证据边界提示，不把待确认内容写成事实。
11. 只返回一个 JSON 对象，不要 Markdown、解释、代码围栏或额外文字。
12. evidence 中每一项的 proof、diagnosis 和 suggestion 必须围绕同一条 JD 要求、同一份简历证据得出，调整建议不得跳到其他能力或凭空新增项目。
13. 评分前先区分硬性门槛、核心要求与加分要求：加分项缺失不得当作硬性缺口重复扣分；已有相邻场景的可迁移证据时，应在“业务场景匹配”中如实反映，而不是按完全无经验处理。
14. 总分区间必须与证据结论一致：82–100 表示核心要求基本都有直接强证据且无关键缺口；64–81 表示核心能力大体可迁移但仍有重要证据或场景缺口；45–63 表示多个核心要求缺少直接证据；0–44 仅用于硬性门槛明显不满足。不要为了迎合用户抬高分数。

通用评分标准：
- 核心能力匹配 30%：JD 最关键的能力要求，简历里是否有直接对应的经历、技能和项目证据。
- 经历证据强度 25%：简历里的工作/项目经历是否足以证明候选人真的做过、能做好相关职责。
- 业务场景匹配 15%：过往业务、行业、用户、产品形态或商业模式是否接近目标岗位所在场景。
- 结果量化程度 10%：简历是否说明结果、指标、规模、影响范围，而不是只写负责/参与。
- 角色责任匹配 10%：候选人在经历中是主导、owner、关键参与，还是辅助执行，是否符合岗位期待。
- 关键词与表达 10%：简历表达是否覆盖 JD 关键词和岗位语言，但不能用关键词堆砌替代真实证据。

置信度规则：
- 高：JD 具体、简历解析完整、引用证据充分且可逐字核验。
- 中：JD 或简历有部分模糊，但核心判断仍有证据。
- 低：JD 太泛、简历文本过短/疑似解析缺失、关键判断依赖推断。

JSON 必须严格包含：
{
  "score_model":"通用证据链评分模型 v1",
  "position_type":"识别到的岗位大类，如产品经理/AI 产品经理/运营/增长/商业化/数据分析/通用岗位",
  "confidence_label":"高/中/低",
  "confidence_reason":"为什么本次评分可信或不够可信",
  "dimensions": [
    {"name":"核心能力匹配","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"},
    {"name":"经历证据强度","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"},
    {"name":"业务场景匹配","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"},
    {"name":"结果量化程度","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"},
    {"name":"角色责任匹配","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"},
    {"name":"关键词与表达","value":0到100整数,"note":"评分依据","bonus_evidence":["加分证据"],"deduction_reason":"扣分原因","remedy":"补救建议"}
  ],
  "summary":"申请建议的核心判断，2到3句",
  "evidence":[{"req":"岗位要求","priority":"硬性门槛/核心要求/能力门槛/加分要求","jdQuote":"JD原文片段","proof":"匹配判断","resumeQuote":"简历中真实存在的原文片段；完全没有相关经历时写无相关证据","level":"直接证据/可迁移证据/无相关证据","confidence":0到100整数,"diagnosis":"以 已匹配/表达缺口/证据不足/真实缺口/待确认 开头，再说明原因","suggestion":"必须同时结合该条JD原文和简历证据，指出具体改哪句话、补哪项事实或如何改写；无法给出具体动作时留空"}],
  "gaps":[{"level":"高/中/低","title":"缺口标题","desc":"具体说明"}],
  "actions":[{"title":"行动标题","desc":"可执行步骤","effort":"预计时间"}],
  "rewrites":[{"before":"简历逐字原文","after":"基于同一事实的建议表达"}],
  "questions":["面试高风险追问"]
}

数量要求：evidence 固定 4 项，gaps 3 项，actions 3 项，rewrites 固定 3 项，questions 固定 5 项。改写建议按重要性排序，只选择对目标岗位最关键的三条；面试问题覆盖核心能力、项目深挖、数据结果、协作推进与岗位动机，并直接使用面试官提问口吻，不要出现“如果面试官追问”“你会如何说明”之类的转述句式。内容应具体、完整，避免为了简短省略判断依据、证据边界或行动步骤。"""


def build_user_prompt(request: AnalyzeRequest) -> str:
    return """目标岗位：{target}
目标公司：{company}
事业部/业务方向：{business_unit}

【简历原文】
{resume}

【JD 原文】
{jd}

请按系统规定的 JSON 结构完成分析。公司与事业部用于判断业务语境和岗位命名，不得替代简历证据。""".format(
        target=request.target,
        company=request.company or "未提供",
        business_unit=request.business_unit or "未提供",
        resume=request.resume,
        jd=request.jd,
    )


def extract_json(content: str) -> Dict[str, Any]:
    cleaned = content.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("model response does not contain a JSON object")
    parsed = json.loads(cleaned[start:end + 1])
    if not isinstance(parsed, dict):
        raise ValueError("model response JSON must be an object")
    return parsed


def normalize_dimensions(raw_dimensions: List[Dimension]) -> List[Dimension]:
    by_name = {item.name: item for item in raw_dimensions}
    normalized: List[Dimension] = []
    for name, weight in DIMENSION_WEIGHTS.items():
        source = by_name.get(name)
        if source is None:
            raise ValueError("missing score dimension: {0}".format(name))
        stable_value = max(0, min(100, int(round(source.value / 5.0) * 5)))
        normalized.append(
            Dimension(
                name=name,
                weight=weight,
                value=stable_value,
                note=source.note,
                points=round(stable_value * weight / 100.0, 1),
                bonus_evidence=source.bonus_evidence[:3],
                deduction_reason=source.deduction_reason,
                remedy=source.remedy,
            )
        )
    return normalized


def grounded_excerpt(quote: str, source: str) -> Optional[str]:
    if quote in ("无直接证据", "无相关证据"):
        return None
    normalized_source = re.sub(r"\s+", "", source)
    candidates = [quote] + re.split(r"[；;。！？!?，,\n]", quote)
    for candidate in sorted(candidates, key=len, reverse=True):
        cleaned = candidate.strip().strip("“”\"'")
        normalized = re.sub(r"\s+", "", cleaned)
        if len(normalized) >= 4 and normalized in normalized_source:
            return cleaned
    return None


def quote_is_grounded(quote: str, source: str) -> bool:
    return grounded_excerpt(quote, source) is not None


def normalize_result(payload: AnalysisPayload, request: AnalyzeRequest) -> AnalysisResult:
    dimensions = normalize_dimensions(payload.dimensions)
    score = int(round(sum(item.points for item in dimensions)))

    checked_evidence: List[Evidence] = []
    for item in payload.evidence[:6]:
        jd_excerpt = grounded_excerpt(item.jdQuote, request.jd)
        resume_excerpt = grounded_excerpt(item.resumeQuote, request.resume)
        if jd_excerpt:
            item.jdQuote = jd_excerpt
        if resume_excerpt:
            item.resumeQuote = resume_excerpt
            if item.level == "无相关证据":
                item.level = "可迁移证据"
        else:
            item.resumeQuote = "无相关证据"
            item.level = "无相关证据"
            item.confidence = min(item.confidence, 55)
            if not item.diagnosis.startswith(("证据不足", "真实缺口")):
                item.diagnosis = "证据不足：" + item.diagnosis
        if not jd_excerpt:
            item.confidence = min(item.confidence, 55)
        checked_evidence.append(item)

    if score >= 82:
        action = "优先投递"
        priority = "A"
    elif score >= 74:
        action = "建议投递"
        priority = "A-"
    elif score >= 64:
        action = "评估后投递"
        priority = "B"
    else:
        action = "暂缓投递"
        priority = "C"

    context_parts = [request.company, request.business_unit, request.target]
    jd_name = " · ".join([item for item in context_parts if item])

    return AnalysisResult(
        target=request.target,
        resumeName="当前上传简历",
        jdName="{0}（当前 JD）".format(jd_name or request.target),
        score=score,
        priority=priority,
        action=action,
        model=model_name(),
        dimensions=dimensions,
        summary=payload.summary,
        evidence=checked_evidence,
        gaps=payload.gaps[:3],
        actions=payload.actions[:3],
        rewrites=payload.rewrites[:3],
        questions=payload.questions[:3],
        score_model=payload.score_model,
        position_type=payload.position_type,
        confidence_label=payload.confidence_label if payload.confidence_label in ("高", "中", "低") else "中",
        confidence_reason=payload.confidence_reason,
    )


async def request_analysis(request: AnalyzeRequest) -> AnalysisPayload:
    api_key = get_secret("SILICONFLOW_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="AI service is not configured")

    messages = [
        {"role": "system", "content": build_system_prompt()},
        {"role": "user", "content": build_user_prompt(request)},
    ]
    headers = {
        "Authorization": "Bearer {0}".format(api_key),
        "Content-Type": "application/json",
    }
    base_payload: Dict[str, Any] = {
        "model": model_name(),
        "messages": messages,
        "temperature": 0,
        "max_tokens": 3800,
        "enable_thinking": False,
        "stream": False,
    }

    last_error: Optional[Exception] = None
    attempts = [True]
    for use_json_format in attempts:
        payload = dict(base_payload)
        if use_json_format:
            payload["response_format"] = {"type": "json_object"}
        try:
            timeout = httpx.Timeout(130.0, connect=15.0)
            async with httpx.AsyncClient(timeout=timeout, trust_env=False) as client:
                response = await client.post(
                    "{0}/chat/completions".format(api_base_url().rstrip("/")),
                    headers=headers,
                    json=payload,
                )
            if response.status_code >= 400:
                detail = response.text[:500]
                raise RuntimeError("SiliconFlow error {0}: {1}".format(response.status_code, detail))
            response_data = response.json()
            content = response_data["choices"][0]["message"]["content"]
            return AnalysisPayload.parse_obj(extract_json(content))
        except Exception as error:
            last_error = error
            LOGGER.warning("AI analysis attempt failed: %s", error)

    raise HTTPException(
        status_code=502,
        detail="模型返回内容未通过报告结构校验，请稍后重试。",
    ) from last_error


async def request_screening(request: ScreeningRequest) -> ScreeningResponse:
    api_key = get_secret("SILICONFLOW_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="AI service is not configured")
    jobs = "\n\n".join(
        "[岗位 {0}]\n岗位名：{1}\n公司：{2}\n方向：{3}\nJD：\n{4}".format(
            index, item.title, item.company or "未提供", item.business_unit or "未提供", item.jd
        )
        for index, item in enumerate(request.jds)
    )
    system_prompt = """你是 Eviden 的多岗位筛选引擎。比较同一份简历与多个 JD，帮助用户决定先投哪个岗位。只依据简历与 JD，不得编造。所有岗位使用同一评分口径。输出一个 JSON 对象：{\"results\":[{\"jd_index\":0,\"title\":\"岗位名\",\"company\":\"公司\",\"score\":0,\"priority\":\"高优先级/中优先级/低优先级\",\"verdict\":\"一句话判断\",\"match_points\":[\"最多3项\"],\"gaps\":[\"最多2项\"]}]}。必须覆盖每个输入岗位并按 score 从高到低排序。不要输出 Markdown。"""
    user_prompt = "【简历原文】\n{0}\n\n【候选岗位】\n{1}".format(request.resume, jobs)
    payload = {
        "model": model_name(),
        "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        "temperature": 0,
        "max_tokens": 3200,
        "enable_thinking": False,
        "stream": False,
        "response_format": {"type": "json_object"},
    }
    try:
        timeout = httpx.Timeout(130.0, connect=15.0)
        async with httpx.AsyncClient(timeout=timeout, trust_env=False) as client:
            response = await client.post(
                "{0}/chat/completions".format(api_base_url().rstrip("/")),
                headers={"Authorization": "Bearer {0}".format(api_key), "Content-Type": "application/json"},
                json=payload,
            )
        if response.status_code >= 400:
            raise RuntimeError("SiliconFlow error {0}".format(response.status_code))
        parsed = extract_json(response.json()["choices"][0]["message"]["content"])
        results = [ScreeningMatch.parse_obj(item) for item in parsed.get("results", [])]
        if len(results) != len(request.jds):
            raise ValueError("岗位筛选结果数量不完整")
        return ScreeningResponse(results=sorted(results, key=lambda item: item.score, reverse=True), model=model_name())
    except Exception as error:
        LOGGER.warning("job screening failed: %s", error)
        raise HTTPException(status_code=502, detail="岗位筛选暂时没有完成，请重新生成。") from error


async def request_resume_rewrite(request: ResumeRewriteRequest) -> ResumeRewriteResult:
    api_key = get_secret("SILICONFLOW_API_KEY")
    if not api_key:
        raise HTTPException(status_code=503, detail="AI service is not configured")

    system_prompt = """你是中文求职简历改写顾问，方法来自 /great-resume 与 /make-resume：只基于用户简历里的事实做表达优化，不捏造公司、职位、时间、指标或项目。输出要产品化、克制、可直接复制到简历。严格返回 JSON。"""
    user_prompt = """目标岗位：{target}
目标公司：{company}
事业部/业务方向：{business_unit}

【简历原文】
{resume}

【目标 JD，可为空】
{jd}

请输出 JSON：
{{
  "positioning":"一句话候选人定位，不超过36字",
  "strategy":"改写策略，2句内",
  "bullets":[{{"section":"所属模块，如工作经历/项目经历/技能摘要","before":"简历逐字原文；找不到完整句时引用最接近片段","after":"建议改写，强调动作、场景、结果和个人边界","reason":"为什么这样改","evidence_boundary":"说明没有新增哪些未经证实的信息"}}],
  "gaps":["还需要补充确认的信息，最多5条"],
  "hr_intro":"给 HR 的中文开场白，80字内",
  "cautions":["不能写进简历的风险点，最多4条"]
}}
要求：bullets 4到7条；before 必须逐字来自简历原文并保持足够完整，以便系统将建议准确合并回原简历；after 不得加入原简历没有的量化结果；不要输出完整简历，只输出高价值改写建议；如果 JD 为空，就按目标岗位进行通用产品化改写。""".format(
        target=request.target,
        company=request.company or "未提供",
        business_unit=request.business_unit or "未提供",
        resume=request.resume,
        jd=request.jd or "未提供",
    )
    headers = {
        "Authorization": "Bearer {0}".format(api_key),
        "Content-Type": "application/json",
    }
    payload: Dict[str, Any] = {
        "model": model_name(),
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.15,
        "max_tokens": 2400,
        "enable_thinking": False,
        "stream": False,
        "response_format": {"type": "json_object"},
    }
    try:
        timeout = httpx.Timeout(90.0, connect=15.0)
        async with httpx.AsyncClient(timeout=timeout, trust_env=False) as client:
            response = await client.post(
                "{0}/chat/completions".format(api_base_url().rstrip("/")),
                headers=headers,
                json=payload,
            )
        if response.status_code >= 400:
            raise RuntimeError("SiliconFlow error {0}: {1}".format(response.status_code, response.text[:500]))
        data = extract_json(response.json()["choices"][0]["message"]["content"])
        return ResumeRewriteResult(
            target=request.target,
            positioning=str(data.get("positioning", ""))[:120],
            strategy=str(data.get("strategy", ""))[:500],
            bullets=[RewriteBullet.parse_obj(item) for item in data.get("bullets", [])[:7]],
            gaps=[str(item) for item in data.get("gaps", [])[:5]],
            hr_intro=str(data.get("hr_intro", ""))[:220],
            cautions=[str(item) for item in data.get("cautions", [])[:4]],
            final_resume=str(data.get("final_resume") or request.resume)[:60000],
            model=model_name(),
        )
    except Exception as error:
        LOGGER.warning("AI rewrite attempt failed: %s", error)
        raise HTTPException(status_code=502, detail="简历改写结果未通过结构校验，请稍后重试。") from error


app = FastAPI(
    title="Eviden Analysis API",
    version="0.1.0",
    docs_url="/api/docs",
    redoc_url="/api/redoc",
    openapi_url="/api/openapi.json",
)
WEB_DIR = os.path.join(os.path.dirname(__file__), "web")
ASSETS_DIR = os.path.join(WEB_DIR, "assets")
if os.path.isdir(ASSETS_DIR):
    app.mount("/assets", StaticFiles(directory=ASSETS_DIR), name="assets")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "OPTIONS", "DELETE"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup_handler():
    init_db()


@app.post("/api/auth/register", response_model=AuthResponse)
async def register_handler(request: AuthRequest):
    now = int(time.time())
    password_hash = hash_password(request.password)
    try:
        user_id = create_user(request.phone, password_hash, now)
    except Exception as error:
        if is_integrity_error(error):
            raise HTTPException(status_code=409, detail="这个手机号已经注册，请直接登录")
        raise
    token = create_session(user_id)
    return AuthResponse(token=token, user=UserInfo(id=user_id, phone=request.phone, created_at=now))


@app.post("/api/auth/login", response_model=AuthResponse)
async def login_handler(request: AuthRequest):
    row, token = authenticate_and_create_session(request.phone, request.password)
    if row is None:
        raise HTTPException(status_code=401, detail="手机号或密码不正确")
    return AuthResponse(token=token, user=user_from_row(row))


@app.get("/api/auth/me", response_model=UserInfo)
async def me_handler(user: UserInfo = Depends(current_user)):
    return user


@app.get("/api/reports", response_model=CloudReportsResponse)
async def list_reports_handler(user: UserInfo = Depends(current_user)):
    rows = fetch_all(
        "SELECT id, record_json, updated_at FROM reports WHERE user_id = ? ORDER BY updated_at DESC",
        (user.id,),
    )
    reports = [
        CloudReport(id=str(row["id"]), record=json.loads(str(row["record_json"])), updated_at=int(row["updated_at"]))
        for row in rows
    ]
    return CloudReportsResponse(reports=reports)


@app.post("/api/reports", response_model=CloudReport)
async def save_report_handler(request: CloudReportRequest, user: UserInfo = Depends(current_user)):
    now = int(time.time())
    record_json = json.dumps(request.record, ensure_ascii=False)
    execute_sql(
        """
        INSERT INTO reports(id, user_id, record_json, updated_at) VALUES (?, ?, ?, ?)
        ON CONFLICT(id, user_id) DO UPDATE SET record_json = excluded.record_json, updated_at = excluded.updated_at
        """,
        (request.id, user.id, record_json, now),
    )
    return CloudReport(id=request.id, record=request.record, updated_at=now)


@app.delete("/api/reports/{report_id}")
async def delete_report_handler(report_id: str, user: UserInfo = Depends(current_user)):
    execute_sql("DELETE FROM reports WHERE id = ? AND user_id = ?", (report_id, user.id))
    return {"ok": True}


@app.get("/api/default-resume", response_model=DefaultResumeResponse)
async def get_default_resume_handler(user: UserInfo = Depends(current_user)):
    row = fetch_one(
        """
        SELECT filename, resume_text, char_count, preview, updated_at
        FROM user_resumes
        WHERE user_id = ?
        """,
        (user.id,),
    )
    return DefaultResumeResponse(resume=default_resume_from_row(row) if row is not None else None)


@app.put("/api/default-resume", response_model=DefaultResumeResponse)
async def put_default_resume_handler(request: DefaultResumeRequest, user: UserInfo = Depends(current_user)):
    now = int(time.time())
    char_count = len(request.resume_text)
    preview = request.resume_text[:260]
    execute_sql(
        """
        INSERT INTO user_resumes(user_id, filename, resume_text, char_count, preview, updated_at)
        VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(user_id) DO UPDATE SET
            filename = excluded.filename,
            resume_text = excluded.resume_text,
            char_count = excluded.char_count,
            preview = excluded.preview,
            updated_at = excluded.updated_at
        """,
        (user.id, request.filename, request.resume_text, char_count, preview, now),
    )
    resume = DefaultResume(
        filename=request.filename,
        resume_text=request.resume_text,
        char_count=char_count,
        preview=preview,
        updated_at=now,
    )
    return DefaultResumeResponse(resume=resume)


@app.delete("/api/default-resume", response_model=DeleteResponse)
async def delete_default_resume_handler(user: UserInfo = Depends(current_user)):
    execute_sql("DELETE FROM user_resumes WHERE user_id = ?", (user.id,))
    return DeleteResponse(ok=True)


@app.get("/")
async def frontend_handler():
    index_path = os.path.join(WEB_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    raise HTTPException(status_code=404, detail="frontend is not built")


@app.get("/api/health", response_model=HealthResponse)
async def health_handler():
    return HealthResponse(
        ok=True,
        model=model_name(),
        configured=bool(get_secret("SILICONFLOW_API_KEY")),
        db_backend="supabase" if use_postgres() else "sqlite",
    )


async def parse_resume_content(filename: str, content: bytes) -> ResumeParseResponse:
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="文件过大，请上传 8MB 以内的简历。")
    try:
        if filename.lower().endswith(".pdf"):
            text = await parse_pdf_resume_with_ocr_fallback(content)
        else:
            text = parse_resume_file(filename, content)
    except HTTPException:
        raise
    except Exception as error:
        LOGGER.warning("resume parse failed: %s", error)
        raise HTTPException(status_code=400, detail="简历解析失败，请换一个 PDF / DOCX 文件重试。")
    if len(text) < 80:
        raise HTTPException(status_code=422, detail="解析到的简历内容太少，请确认文件不是扫描图片或空白文档。")
    return ResumeParseResponse(
        filename=filename,
        text=text,
        char_count=len(text),
        preview=text[:260],
    )


@app.post("/api/parse-resume", response_model=ResumeParseResponse)
async def parse_resume_handler(file: UploadFile = File(...)):
    filename = file.filename or "resume"
    return await parse_resume_content(filename, await file.read())


@app.post("/api/parse-resume-json", response_model=ResumeParseResponse)
async def parse_resume_json_handler(request: ResumeParseJsonRequest):
    try:
        content = base64.b64decode(request.content_base64, validate=True)
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="简历文件读取失败，请重新选择文件。")
    return await parse_resume_content(request.filename, content)


async def process_analysis_job(job_id: str, request: AnalyzeRequest) -> None:
    now = int(time.time())
    with _ANALYSIS_JOBS_LOCK:
        _ANALYSIS_JOBS[job_id].update(status="processing", updated_at=now)
    try:
        payload = await request_analysis(request)
        result = normalize_result(payload, request)
        now = int(time.time())
        with _ANALYSIS_JOBS_LOCK:
            _ANALYSIS_JOBS[job_id].update(status="completed", result=result.dict(), error="", updated_at=now)
    except Exception as error:
        LOGGER.exception("analysis job failed: %s", job_id)
        detail = getattr(error, "detail", "") or "报告生成失败，请重新生成。"
        now = int(time.time())
        with _ANALYSIS_JOBS_LOCK:
            _ANALYSIS_JOBS[job_id].update(status="failed", error=str(detail)[:500], updated_at=now)


def start_analysis_job(job_id: str, request: AnalyzeRequest) -> None:
    asyncio.run(process_analysis_job(job_id, request))


def analysis_job_response(job_id: str, job: Dict[str, Any]) -> AnalysisJobStatus:
    result = job.get("result")
    return AnalysisJobStatus(
        job_id=job_id,
        status=str(job["status"]),
        result=AnalysisResult.parse_obj(result) if result else None,
        error=str(job.get("error") or ""),
        created_at=int(job["created_at"]),
        updated_at=int(job["updated_at"]),
    )


@app.post("/api/analysis-jobs", response_model=AnalysisJobCreated)
async def create_analysis_job_handler(request: AnalyzeRequest):
    job_id = "analysis-{0}".format(uuid.uuid4().hex)
    now = int(time.time())
    with _ANALYSIS_JOBS_LOCK:
        expired_before = now - 60 * 60
        for old_job_id in [key for key, value in _ANALYSIS_JOBS.items() if value["updated_at"] < expired_before]:
            _ANALYSIS_JOBS.pop(old_job_id, None)
        _ANALYSIS_JOBS[job_id] = {
            "status": "queued",
            "result": None,
            "error": "",
            "created_at": now,
            "updated_at": now,
        }
    threading.Thread(target=start_analysis_job, args=(job_id, request), daemon=True).start()
    return AnalysisJobCreated(job_id=job_id, status="queued")


@app.get("/api/analysis-jobs/{job_id}", response_model=AnalysisJobStatus)
async def get_analysis_job_handler(job_id: str):
    with _ANALYSIS_JOBS_LOCK:
        job = dict(_ANALYSIS_JOBS.get(job_id) or {})
    if not job:
        raise HTTPException(status_code=404, detail="生成任务不存在或已失效。")
    return analysis_job_response(job_id, job)


async def process_rewrite_job(job_id: str, request: ResumeRewriteRequest) -> None:
    now = int(time.time())
    with _ANALYSIS_JOBS_LOCK:
        _ANALYSIS_JOBS[job_id].update(status="processing", updated_at=now)
    try:
        result = await request_resume_rewrite(request)
        now = int(time.time())
        with _ANALYSIS_JOBS_LOCK:
            _ANALYSIS_JOBS[job_id].update(status="completed", result=result.dict(), error="", updated_at=now)
    except Exception as error:
        LOGGER.exception("resume rewrite job failed: %s", job_id)
        detail = getattr(error, "detail", "") or "简历优化失败，请重新生成。"
        now = int(time.time())
        with _ANALYSIS_JOBS_LOCK:
            _ANALYSIS_JOBS[job_id].update(status="failed", error=str(detail)[:500], updated_at=now)


def start_rewrite_job(job_id: str, request: ResumeRewriteRequest) -> None:
    asyncio.run(process_rewrite_job(job_id, request))


@app.post("/api/rewrite-jobs", response_model=AnalysisJobCreated)
async def create_rewrite_job_handler(request: ResumeRewriteRequest):
    job_id = "rewrite-{0}".format(uuid.uuid4().hex)
    now = int(time.time())
    with _ANALYSIS_JOBS_LOCK:
        expired_before = now - 60 * 60
        for old_job_id in [key for key, value in _ANALYSIS_JOBS.items() if value["updated_at"] < expired_before]:
            _ANALYSIS_JOBS.pop(old_job_id, None)
        _ANALYSIS_JOBS[job_id] = {
            "status": "queued",
            "result": None,
            "error": "",
            "created_at": now,
            "updated_at": now,
        }
    threading.Thread(target=start_rewrite_job, args=(job_id, request), daemon=True).start()
    return AnalysisJobCreated(job_id=job_id, status="queued")


@app.get("/api/rewrite-jobs/{job_id}", response_model=ResumeRewriteJobStatus)
async def get_rewrite_job_handler(job_id: str):
    with _ANALYSIS_JOBS_LOCK:
        job = dict(_ANALYSIS_JOBS.get(job_id) or {})
    if not job or not job_id.startswith("rewrite-"):
        raise HTTPException(status_code=404, detail="优化任务不存在或已失效。")
    result = job.get("result")
    return ResumeRewriteJobStatus(
        job_id=job_id,
        status=str(job["status"]),
        result=ResumeRewriteResult.parse_obj(result) if result else None,
        error=str(job.get("error") or ""),
        created_at=int(job["created_at"]),
        updated_at=int(job["updated_at"]),
    )


@app.post("/api/screen-jobs", response_model=ScreeningResponse)
async def screen_jobs_handler(request: ScreeningRequest):
    return await request_screening(request)


@app.post("/api/analyze", response_model=AnalysisResult)
async def analyze_handler(request: AnalyzeRequest):
    payload = await request_analysis(request)
    return normalize_result(payload, request)


@app.post("/api/rewrite-resume", response_model=ResumeRewriteResult)
async def rewrite_resume_handler(request: ResumeRewriteRequest):
    return await request_resume_rewrite(request)


# ---------------------------DO NOT EDIT CODE BELOW THIS LINE---------------------------------
# This is the entry point for the FastAPI application.
if __name__ == "__main__":
    port = int(os.environ.get("_BYTEFAAS_RUNTIME_PORT", 8000))
    config = uvicorn.Config("main:app", port=port, log_level="info", host=None)
    server = uvicorn.Server(config)
    server.run()
# --------------------------------------------------------------------------------------------
